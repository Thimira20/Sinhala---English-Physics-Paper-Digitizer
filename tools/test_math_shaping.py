"""
Test script to compare different OMML XML representations for Sinhala text shaping.
"""
from pathlib import Path
import sys
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from core.word_com import convert_to_pdf
try:
    import pymupdf as fitz
except ImportError:
    import fitz

def main():
    doc = Document(str(PROJECT_ROOT / "assets" / "reference.docx"))

    # 1. Option A (Current: m:r)
    p1 = doc.add_paragraph("Option A (Current m:r):")
    xml_a = (
        f'<m:oMathPara {nsdecls("m", "w")}>'
        f'<m:oMath>'
        f'<m:r><m:rPr><m:nor/><m:sty m:val="p"/></m:rPr>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<m:t>භාගික දෝෂය</m:t></m:r>'
        f'<m:r><m:t> = </m:t></m:r>'
        f'<m:f><m:num><m:r><m:rPr><m:nor/><m:sty m:val="p"/></m:rPr>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<m:t>දෝෂය</m:t></m:r></m:num>'
        f'<m:den><m:r><m:rPr><m:nor/><m:sty m:val="p"/></m:rPr>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<m:t>පාඨාංකය</m:t></m:r></m:den></m:f>'
        f'</m:oMath></m:oMathPara>'
    )
    p1._element.append(parse_xml(xml_a))

    # 2. Option B (w:r inside m:oMath - Word text run)
    p2 = doc.add_paragraph("Option B (w:r inside oMath):")
    xml_b = (
        f'<m:oMathPara {nsdecls("m", "w")}>'
        f'<m:oMath>'
        f'<w:r>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<w:t>භාගික දෝෂය</w:t></w:r>'
        f'<m:r><m:t> = </m:t></m:r>'
        f'<m:f><m:num><w:r>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<w:t>දෝෂය</w:t></w:r></m:num>'
        f'<m:den><w:r>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<w:t>පාඨාංකය</w:t></w:r></m:den></m:f>'
        f'</m:oMath></m:oMathPara>'
    )
    p2._element.append(parse_xml(xml_b))

    # 3. Option C (Iskoola Pota in m:r)
    p3 = doc.add_paragraph("Option C (Iskoola Pota in m:r):")
    xml_c = (
        f'<m:oMathPara {nsdecls("m", "w")}>'
        f'<m:oMath>'
        f'<m:r><m:rPr><m:nor/><m:sty m:val="p"/></m:rPr>'
        f'<w:rPr><w:rFonts w:ascii="Iskoola Pota" w:hAnsi="Iskoola Pota" w:cs="Iskoola Pota"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<m:t>භාගික දෝෂය</m:t></m:r>'
        f'<m:r><m:t> = </m:t></m:r>'
        f'<m:f><m:num><m:r><m:rPr><m:nor/><m:sty m:val="p"/></m:rPr>'
        f'<w:rPr><w:rFonts w:ascii="Iskoola Pota" w:hAnsi="Iskoola Pota" w:cs="Iskoola Pota"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<m:t>දෝෂය</m:t></m:r></m:num>'
        f'<m:den><m:r><m:rPr><m:nor/><m:sty m:val="p"/></m:rPr>'
        f'<w:rPr><w:rFonts w:ascii="Iskoola Pota" w:hAnsi="Iskoola Pota" w:cs="Iskoola Pota"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<m:t>පාඨාංකය</m:t></m:r></m:den></m:f>'
        f'</m:oMath></m:oMathPara>'
    )
    p3._element.append(parse_xml(xml_c))

    # 4. Option D (Leading text outside equation)
    p4 = doc.add_paragraph()
    r = p4.add_run("Option D: භාගික දෝෂය = ")
    r.font.name = "Noto Sans Sinhala"
    xml_d = (
        f'<m:oMath {nsdecls("m", "w")}>'
        f'<m:f><m:num><w:r>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<w:t>දෝෂය</w:t></w:r></m:num>'
        f'<m:den><w:r>'
        f'<w:rPr><w:rFonts w:ascii="Noto Sans Sinhala" w:hAnsi="Noto Sans Sinhala" w:cs="Noto Sans Sinhala"/><w:lang w:bidi="si-LK"/></w:rPr>'
        f'<w:t>පාඨාංකය</w:t></w:r></m:den></m:f>'
        f'</m:oMath>'
    )
    p4._element.append(parse_xml(xml_d))

    test_docx = PROJECT_ROOT / "workspace" / "exports" / "test_math_options.docx"
    test_pdf = PROJECT_ROOT / "workspace" / "exports" / "test_math_options.pdf"
    doc.save(str(test_docx))
    print(f"Saved: {test_docx}")
    convert_to_pdf(test_docx, test_pdf)
    print(f"PDF generated: {test_pdf}")

    # Render to PNG
    pdf_doc = fitz.open(str(test_pdf))
    page = pdf_doc[0]
    pix = page.get_pixmap(dpi=200)
    out_png = PROJECT_ROOT / "workspace" / "exports" / "test_math_options.png"
    pix.save(str(out_png))
    print(f"Rendered image: {out_png}")

if __name__ == "__main__":
    main()
