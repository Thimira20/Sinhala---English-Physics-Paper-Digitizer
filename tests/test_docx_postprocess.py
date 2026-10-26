"""
Unit tests for core.docx_postprocess (SPEC P3-T04 & P3-T05).
Validates OpenXML manipulation for Sinhala fonts, bold CS, OMML math child ordering,
and MCQ table restyling.
"""
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from core.docx_postprocess import (
    process_w_run,
    process_omml_run,
    process_mcq_tables,
    fix_sinhala_and_styles,
    RPR_CHILD_ORDER
)

def test_process_w_run_sinhala():
    # Run with Sinhala text and bold
    r_xml = f'''
        <w:r {nsdecls("w")}>
            <w:rPr>
                <w:b/>
                <w:sz w:val="24"/>
            </w:rPr>
            <w:t>ප්‍රශ්නය 01</w:t>
        </w:r>
    '''
    r_el = parse_xml(r_xml)
    process_w_run(r_el, sinhala_font="Noto Sans Sinhala", latin_font="Noto Sans")

    rpr = r_el.find(qn("w:rPr"))
    assert rpr is not None

    # Check fonts
    rfonts = rpr.find(qn("w:rFonts"))
    assert rfonts is not None
    assert rfonts.get(qn("w:cs")) == "Noto Sans Sinhala"
    assert rfonts.get(qn("w:ascii")) == "Noto Sans"

    # Check mirrored properties
    assert rpr.find(qn("w:bCs")) is not None
    sz_cs = rpr.find(qn("w:szCs"))
    assert sz_cs is not None and sz_cs.get(qn("w:val")) == "24"

    # Check bidi language
    lang = rpr.find(qn("w:lang"))
    assert lang is not None and lang.get(qn("w:bidi")) == "si-LK"

    # Validate child order according to RPR_CHILD_ORDER
    tag_indices = [RPR_CHILD_ORDER.index(child.tag.split("}")[-1]) for child in rpr if child.tag.split("}")[-1] in RPR_CHILD_ORDER]
    assert tag_indices == sorted(tag_indices)

def test_process_omml_run_with_sinhala():
    # OMML run with Sinhala inside m:t
    mr_xml = f'''
        <m:r {nsdecls("m", "w")}>
            <m:t>වේගය</m:t>
        </m:r>
    '''
    mr_el = parse_xml(mr_xml)
    process_omml_run(mr_el, sinhala_font="Noto Sans Sinhala")

    # Verify child order: m:rPr -> w:rPr -> m:t
    children_tags = [c.tag.split("}")[-1] for c in mr_el]
    assert children_tags == ["rPr", "rPr", "t"]
    assert mr_el[0].tag == qn("m:rPr")
    assert mr_el[1].tag == qn("w:rPr")
    assert mr_el[2].tag == qn("m:t")

    # Verify m:nor and m:sty inside m:rPr
    mrpr = mr_el[0]
    assert mrpr.find(qn("m:nor")) is not None
    assert mrpr.find(qn("m:sty")) is not None

    # Verify w:rFonts inside w:rPr
    wrpr = mr_el[1]
    assert wrpr.find(qn("w:rFonts")) is not None
    assert wrpr.find(qn("w:lang")) is not None

def test_process_omml_run_latin_math_untouched():
    # OMML run without Sinhala (e.g. x^2 + y^2)
    mr_xml = f'''
        <m:r {nsdecls("m")}>
            <m:t>x^2</m:t>
        </m:r>
    '''
    mr_el = parse_xml(mr_xml)
    process_omml_run(mr_el, sinhala_font="Noto Sans Sinhala")
    # Untouched: no w:rPr or m:nor added
    assert mr_el.find(qn("w:rPr")) is None
    assert mr_el.find(qn("m:rPr")) is None

def test_mcq_tables_postprocessing():
    doc_xml = f'''
        <w:document {nsdecls("w")}>
            <w:body>
                <w:p>
                    <w:pPr>
                        <w:pStyle w:val="MCQMarker"/>
                    </w:pPr>
                </w:p>
                <w:tbl>
                    <w:tblPr>
                        <w:tblStyle w:val="TableGrid"/>
                    </w:tblPr>
                    <w:tr>
                        <w:tc><w:p><w:t> </w:t></w:p></w:tc>
                        <w:tc><w:p><w:t></w:t></w:p></w:tc>
                    </w:tr>
                    <w:tr>
                        <w:tc><w:p><w:t>(1) 10 N</w:t></w:p></w:tc>
                        <w:tc><w:p><w:t>(2) 20 N</w:t></w:p></w:tc>
                    </w:tr>
                </w:tbl>
            </w:body>
        </w:document>
    '''
    doc_el = parse_xml(doc_xml)
    process_mcq_tables(doc_el)

    body = doc_el.find(qn("w:body"))
    # MCQMarker paragraph must be deleted
    pstyles = [p.find(f".//{qn('w:pStyle')}") for p in body.findall(qn("w:p"))]
    assert not any(ps is not None and ps.get(qn("w:val")) == "MCQMarker" for ps in pstyles)

    # Table style must be updated to MCQTable
    tbl = body.find(qn("w:tbl"))
    assert tbl is not None
    tbl_style = tbl.find(f".//{qn('w:tblStyle')}")
    assert tbl_style is not None and tbl_style.get(qn("w:val")) == "MCQTable"

    # Empty header row must be removed (only 1 data row remains)
    rows = tbl.findall(qn("w:tr"))
    assert len(rows) == 1
    cell_text = "".join([t.text for t in rows[0].findall(f".//{qn('w:t')}") if t.text])
    assert "(1) 10 N" in cell_text

def test_full_docx_postprocess(tmp_path):
    doc_path = tmp_path / "test.docx"
    doc = Document()
    p = doc.add_paragraph("මෙය පරීක්ෂණ ඡේදයකි.")
    doc.save(str(doc_path))

    fix_sinhala_and_styles(doc_path, sinhala_font="Noto Sans Sinhala", latin_font="Noto Sans")

    # Reopen and verify
    reopened = Document(str(doc_path))
    p_el = reopened.paragraphs[0]._element
    r_el = p_el.find(qn("w:r"))
    rpr = r_el.find(qn("w:rPr"))
    rfonts = rpr.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:cs")) == "Noto Sans Sinhala"
